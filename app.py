# ╔══════════════════════════════════════════════════════════════════════╗
# ║          Brand Analysis Tool  —  app_v3.py                         ║
# ║  Section 1 : Discount Analysis  (Matrix + Insights)                ║
# ║  Section 2 : Product Analysis   (4-Quadrant overall view)          ║
# ╚══════════════════════════════════════════════════════════════════════╝

import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

st.set_page_config(page_title="Brand Analysis Tool", layout="wide", page_icon="📊")

# ══════════════════════════════════════════════════════════════════════
#  SHARED HELPERS
# ══════════════════════════════════════════════════════════════════════
MONTH_MAP = {1:'January',2:'February',3:'March',4:'April',5:'May',6:'June',
             7:'July',8:'August',9:'September',10:'October',11:'November',12:'December'}

def parse_month_start(s):
    return pd.to_datetime(str(s).split(' - ')[0].strip())

def make_month_label(s):
    dt = parse_month_start(s)
    return f"{MONTH_MAP[dt.month]} {dt.year}"

def clean_pid(x):
    try:    return str(int(float(x)))
    except: return str(x).strip()

def find_col(df, *kws):
    for kw in kws:
        hits = [c for c in df.columns if kw.lower() in c.lower()]
        if hits: return hits[0]
    return None

def fmt_inr(v):  return f"₹{v:,.0f}"
def fmt_roi(v):  return f"{v:.2f}x"
def fmt_pct(v):  return f"{v*100:.1f}%"

# ══════════════════════════════════════════════════════════════════════
#  SIDEBAR  — navigation + all file uploads
# ══════════════════════════════════════════════════════════════════════
with st.sidebar:
    st.markdown("## 📊 Brand Analysis Tool")
    st.markdown("---")

    page = st.radio("Navigate to", ["📋 Discount Analysis", "🔲 Product Analysis"],
                    label_visibility="collapsed")

    st.markdown("---")

    # ── Section 1 files ───────────────────────────────────────────────
    if page == "📋 Discount Analysis":
        st.markdown("### 📁 Upload Files")

        with st.expander("📌 Required CSV format", expanded=False):
            st.markdown("""
**Meta Spend CSV**
| Column | Description |
|---|---|
| `Product ID` | Unique product identifier (number) |
| `Product Title` | Product name |
| `Month` | Format: `2026-03-01 - 2026-03-31` |
| `Amount spent (INR)` | Ad spend in INR (numeric) |

**Shopify Revenue CSV**
| Column | Description |
|---|---|
| `Product ID` | Same product identifier as Meta |
| `Product title` | Product name |
| `Month` | Same format as Meta |
| `Net sales` | Revenue in INR (numeric) |

**Discount Product List CSV**
| Column | Description |
|---|---|
| `Product ID` | IDs of discounted products |
| `Product Title` | Product name (optional) |
""")

        meta_file     = st.file_uploader("① Meta Spend CSV",        type=['csv'], key='s1_meta')
        shopify_file  = st.file_uploader("② Shopify Revenue CSV",   type=['csv'], key='s1_shop')
        discount_file = st.file_uploader("③ Discount Product List", type=['csv','xlsx'], key='s1_disc')
        brand_name    = st.text_input("Brand Name", value="Brand", key='s1_brand')

        st.markdown("---")
        st.markdown("**🎚 Insight Thresholds**")
        st.caption("High Spend = spend above this % of category average. Low Revenue = revenue below this % of category average.")
        spend_pct = st.slider("High Spend threshold (%)", 50, 300, 100, 10, key='s1_sp',
                              help="e.g. 100 = products spending more than the category average")
        rev_pct   = st.slider("Low Revenue threshold (%)", 10, 200, 100, 10, key='s1_rv',
                              help="e.g. 100 = products earning less than the category average")

        run_s1 = st.button("▶ Run Discount Analysis", type="primary")

    # ── Section 2 files ───────────────────────────────────────────────
    else:
        st.markdown("### 📁 Upload Files")

        with st.expander("📌 Required CSV format", expanded=False):
            st.markdown("""
**Meta Spend CSV**
| Column | Description |
|---|---|
| `Product ID` | Unique product identifier (number) |
| `Product Title` | Product name |
| `Month` | Format: `2026-03-01 - 2026-03-31` |
| `Amount spent (INR)` | Ad spend in INR (numeric) |

**Shopify Revenue CSV**
| Column | Description |
|---|---|
| `Product ID` | Same product identifier as Meta |
| `Product title` | Product name |
| `Month` | Same format as Meta |
| `Net sales` | Revenue in INR (numeric) |

> ℹ️ No Discount List needed here — this section analyses all products.
""")

        s2_meta_file   = st.file_uploader("① Meta Spend CSV",      type=['csv'], key='s2_meta')
        s2_shopify_file= st.file_uploader("② Shopify Revenue CSV", type=['csv'], key='s2_shop')

        st.markdown("---")
        st.markdown("**🎚 Quadrant Thresholds**")
        st.caption("A product is 'High Spend / High Revenue' when its value exceeds X% of the overall product average.")
        s2_spend_pct = st.slider("High Spend threshold (%)", 50, 300, 100, 10, key='s2_sp',
                                 help="100% = above average spend across all products")
        s2_rev_pct   = st.slider("High Revenue threshold (%)", 50, 300, 100, 10, key='s2_rv',
                                 help="100% = above average revenue across all products")

        run_s2 = st.button("▶ Run Product Analysis", type="primary")


# ══════════════════════════════════════════════════════════════════════
#  SECTION 1 — DISCOUNT ANALYSIS ENGINE
# ══════════════════════════════════════════════════════════════════════
def run_discount_analysis(meta, shopify, discount, spend_pct_thresh, rev_pct_thresh):
    meta_spend_col  = find_col(meta,    'spent','spend','amount')
    shopify_rev_col = find_col(shopify, 'sales','revenue','net sales')
    meta_title_col  = find_col(meta,    'product title','title','name')
    shop_title_col  = find_col(shopify, 'product title','title','name')

    if not meta_spend_col:
        raise ValueError("Cannot detect spend column in Meta CSV.")
    if not shopify_rev_col:
        raise ValueError("Cannot detect revenue column in Shopify CSV.")

    meta    = meta.copy();    shopify = shopify.copy();    discount = discount.copy()
    meta['_pid']     = meta['Product ID'].apply(clean_pid)
    shopify['_pid']  = shopify['Product ID'].apply(clean_pid)
    discount['_pid'] = discount['Product ID'].dropna().apply(clean_pid)

    # Title map: Meta overwrites Shopify
    title_map = {}
    if shop_title_col:
        for _, r in shopify.drop_duplicates('_pid').iterrows():
            title_map[r['_pid']] = r[shop_title_col]
    if meta_title_col:
        for _, r in meta.drop_duplicates('_pid').iterrows():
            title_map[r['_pid']] = r[meta_title_col]

    meta['_spend']  = pd.to_numeric(meta[meta_spend_col],  errors='coerce').fillna(0)
    shopify['_rev'] = pd.to_numeric(shopify[shopify_rev_col], errors='coerce').fillna(0)

    meta_g    = meta.groupby(['_pid','Month'])['_spend'].sum().reset_index()
    shopify_g = shopify.groupby(['_pid','Month'])['_rev'].sum().reset_index()

    merged = pd.merge(meta_g, shopify_g, on=['_pid','Month'], how='outer').fillna(0)
    merged.columns = ['Product ID','Month','Spend','Revenue']

    disc_ids = set(discount['_pid'].dropna())
    merged['Is_Discounted'] = merged['Product ID'].isin(disc_ids)

    months_raw     = sorted(merged['Month'].dropna().unique(), key=parse_month_start)
    merged['Month_Label'] = merged['Month'].apply(make_month_label)
    months_ordered = [make_month_label(m) for m in months_raw]

    # Summary matrix
    results = []
    for month in months_ordered:
        md = merged[merged['Month_Label']==month]
        ts, tr = md['Spend'].sum(), md['Revenue'].sum()
        for is_d, cat in [(True,'Discounted'),(False,'Non-Discounted')]:
            g = md[md['Is_Discounted']==is_d]
            sp, rv = g['Spend'].sum(), g['Revenue'].sum()
            results.append({'Month':month,'Category':cat,
                'Spend':round(sp,2),'Revenue':round(rv,2),
                'Spend_Pct':  round(sp/ts,4) if ts else 0,
                'Revenue_Pct':round(rv/tr,4) if tr else 0,
                'ROI':        round(rv/sp,4) if sp else 0})

    # Insights
    insights = {}
    for month in months_ordered:
        md = merged[(merged['Month_Label']==month) & (merged['Spend']>0)].copy()
        md['ROI'] = (md['Revenue']/md['Spend']).round(4)
        for is_d, cat in [(True,'Discounted'),(False,'Non-Discounted')]:
            g = md[md['Is_Discounted']==is_d].copy()
            if g.empty:
                insights[(month,cat)] = {'hslr':pd.DataFrame(),'lshr':pd.DataFrame()}
                continue
            avg_sp = g['Spend'].mean()
            avg_rv = g['Revenue'].mean()
            sp_cut = avg_sp * spend_pct_thresh / 100
            rv_cut = avg_rv * rev_pct_thresh   / 100
            hslr = g[(g['Spend']>=sp_cut) & (g['Revenue']<=rv_cut)].copy()
            lshr = g[(g['Spend']< sp_cut) & (g['Revenue']> rv_cut)].copy()
            for df in [hslr, lshr]:
                df['Product Title'] = df['Product ID'].map(title_map).fillna('Unknown')
            cols = ['Product ID','Product Title','Spend','Revenue','ROI']
            insights[(month,cat)] = {
                'hslr': hslr.sort_values('Spend',   ascending=False)[cols].reset_index(drop=True),
                'lshr': lshr.sort_values('Revenue', ascending=False)[cols].reset_index(drop=True),
                'sp_cut': sp_cut, 'rv_cut': rv_cut,
            }

    return pd.DataFrame(results), months_ordered, merged, insights, title_map


# ── Excel builder for Section 1 ───────────────────────────────────────
def build_s1_excel(results_df, months_ordered, insights):
    wb  = Workbook()
    THIN = Side(style='thin', color='CCCCCC')

    def bdr():
        return Border(left=THIN,right=THIN,top=THIN,bottom=THIN)
    def hdr(cell, text, bg, fg='FFFFFF', sz=10):
        cell.value=text; cell.font=Font(bold=True,color=fg,name='Calibri',size=sz)
        cell.fill=PatternFill('solid',start_color=bg)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=bdr()
    def val(cell, v, fmt, bg=None):
        cell.value=v; cell.number_format=fmt
        cell.font=Font(name='Calibri',size=10)
        cell.alignment=Alignment(horizontal='center',vertical='center')
        cell.border=bdr()
        if bg: cell.fill=PatternFill('solid',start_color=bg)
    def lbl(cell, text, bg='F2F2F2'):
        cell.value=text; cell.font=Font(bold=True,name='Calibri',size=10)
        cell.fill=PatternFill('solid',start_color=bg)
        cell.alignment=Alignment(horizontal='left',vertical='center')
        cell.border=bdr()
    def tot(cell, v, fmt, bg):
        cell.value=v; cell.number_format=fmt
        cell.font=Font(bold=True,name='Calibri',size=10,color='FFFFFF')
        cell.fill=PatternFill('solid',start_color=bg)
        cell.alignment=Alignment(horizontal='center',vertical='center')
        cell.border=bdr()

    DISC_BG='D6E4F0'; NDISC_BG='E8F5E9'

    # ── Sheet 1: Summary Matrix ───────────────────────────────────────
    ws = wb.active; ws.title="Summary Matrix"
    ws.sheet_view.showGridLines=False
    ws.column_dimensions['A'].width=26
    ws.row_dimensions[1].height=28; ws.row_dimensions[2].height=22; ws.row_dimensions[3].height=20

    for i, month in enumerate(months_ordered):
        dc=get_column_letter(2+i*2); nc=get_column_letter(3+i*2)
        ws.column_dimensions[dc].width=22; ws.column_dimensions[nc].width=22
        ws.merge_cells(f'{dc}1:{nc}1')
        hdr(ws[f'{dc}1'], month.upper(), '1F3864', sz=11)
        hdr(ws[f'{dc}2'], 'Discounted',     '2E75B6')
        hdr(ws[f'{nc}2'], 'Non-Discounted', '375623')
        hdr(ws[f'{dc}3'], 'Value', '4472C4', sz=9)
        hdr(ws[f'{nc}3'], 'Value', '4472C4', sz=9)

    ws.merge_cells('A1:A3'); hdr(ws['A1'],'METRIC','1F3864',sz=11)
    ws.freeze_panes='A4'

    metrics=[('Total Spend (INR)','Spend','#,##0.00'),
             ('Total Revenue (INR)','Revenue','#,##0.00'),
             ('Spend %','Spend_Pct','0.0%'),
             ('Revenue %','Revenue_Pct','0.0%'),
             ('ROI','ROI','0.00"x"')]

    for r,(label,key,fmt) in enumerate(metrics,start=4):
        ws.row_dimensions[r].height=20; lbl(ws[f'A{r}'],label)
        for i,month in enumerate(months_ordered):
            dc=get_column_letter(2+i*2); nc=get_column_letter(3+i*2)
            dr=results_df[(results_df['Month']==month)&(results_df['Category']=='Discounted')].iloc[0]
            nr=results_df[(results_df['Month']==month)&(results_df['Category']=='Non-Discounted')].iloc[0]
            val(ws[f'{dc}{r}'],dr[key],fmt,DISC_BG)
            val(ws[f'{nc}{r}'],nr[key],fmt,NDISC_BG)

    # ── Sheet 2: Product Insights ─────────────────────────────────────
    pi=wb.create_sheet("Product Insights"); pi.sheet_view.showGridLines=False
    SEC_COLORS={
        ('Discounted',    'hslr'):('C0392B','FDEDEC','E74C3C'),
        ('Discounted',    'lshr'):('1A5276','D6EAF8','2874A6'),
        ('Non-Discounted','hslr'):('784212','FDEBD0','E67E22'),
        ('Non-Discounted','lshr'):('145A32','D5F5E3','1E8449'),
    }
    SEC_LABELS={'hslr':'⚠ High Spend · Low Revenue  (Underperforming — pause/review)',
                'lshr':'✅ Low Spend · High Revenue  (Efficient — consider scaling)'}
    FMTS=['@','@','#,##0.00','#,##0.00','0.00"x"']
    pi.column_dimensions['A'].width=20; pi.column_dimensions['B'].width=48
    pi.column_dimensions['C'].width=16; pi.column_dimensions['D'].width=16
    pi.column_dimensions['E'].width=10; pi.column_dimensions['F'].width=14

    cr=1
    for month in months_ordered:
        pi.merge_cells(f'A{cr}:F{cr}')
        c=pi.cell(row=cr,column=1,value=f'  {month.upper()}')
        c.font=Font(bold=True,color='FFFFFF',name='Calibri',size=13)
        c.fill=PatternFill('solid',start_color='0D1F3C')
        c.alignment=Alignment(vertical='center'); pi.row_dimensions[cr].height=30; cr+=1

        for cat in ['Discounted','Non-Discounted']:
            ins=insights.get((month,cat),{})
            for itype in ['hslr','lshr']:
                hdr_bg,row_bg,sec_bg=SEC_COLORS[(cat,itype)]
                pi.merge_cells(f'A{cr}:F{cr}')
                sc=pi.cell(row=cr,column=1,value=f'  {cat}  ·  {SEC_LABELS[itype]}')
                sc.font=Font(bold=True,color='FFFFFF',name='Calibri',size=10)
                sc.fill=PatternFill('solid',start_color=sec_bg)
                sc.alignment=Alignment(vertical='center'); pi.row_dimensions[cr].height=22; cr+=1

                df_sec=ins.get(itype,pd.DataFrame())
                if df_sec.empty:
                    pi.merge_cells(f'A{cr}:F{cr}')
                    nc2=pi.cell(row=cr,column=1,value='  No products match this criteria.')
                    nc2.font=Font(italic=True,color='888888',name='Calibri',size=10)
                    nc2.alignment=Alignment(vertical='center'); pi.row_dimensions[cr].height=18; cr+=1
                else:
                    # column headers
                    for ci,ltext in enumerate(['Product ID','Product Title','Spend (INR)','Revenue (INR)','ROI'],1):
                        hdr(pi.cell(row=cr,column=ci),ltext,hdr_bg,sz=9)
                    pi.row_dimensions[cr].height=18; cr+=1
                    # data rows
                    for ri,row in enumerate(df_sec.itertuples(index=False)):
                        bg=row_bg if ri%2==0 else 'FDFEFE'
                        for ci,(cv,fmt) in enumerate(zip(row,FMTS),1):
                            c2=pi.cell(row=cr,column=ci,value=cv)
                            c2.number_format=fmt; c2.font=Font(name='Calibri',size=10)
                            c2.fill=PatternFill('solid',start_color=bg)
                            c2.alignment=Alignment(vertical='center',wrap_text=(ci==2))
                            c2.border=Border(left=Side(style='thin',color='DDDDDD'),
                                             right=Side(style='thin',color='DDDDDD'),
                                             top=Side(style='thin',color='DDDDDD'),
                                             bottom=Side(style='thin',color='DDDDDD'))
                        pi.row_dimensions[cr].height=16; cr+=1
                    # TOTALS ROW
                    t_sp=df_sec['Spend'].sum(); t_rv=df_sec['Revenue'].sum()
                    t_roi=round(t_rv/t_sp,4) if t_sp else 0
                    tot_data=[f'{len(df_sec)} products','TOTAL',t_sp,t_rv,t_roi]
                    tot_fmts=['@','@','#,##0.00','#,##0.00','0.00"x"']
                    for ci,(cv,fmt) in enumerate(zip(tot_data,tot_fmts),1):
                        tot(pi.cell(row=cr,column=ci),cv,fmt,hdr_bg)
                    pi.row_dimensions[cr].height=18; cr+=1

                cr+=1  # gap
        cr+=2  # month gap

    pi.freeze_panes='A2'
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════
#  SECTION 2 — PRODUCT ANALYSIS ENGINE
# ══════════════════════════════════════════════════════════════════════
def run_product_analysis(meta, shopify, spend_pct_thresh, rev_pct_thresh):
    meta_spend_col  = find_col(meta,    'spent','spend','amount')
    shopify_rev_col = find_col(shopify, 'sales','revenue','net sales')
    meta_title_col  = find_col(meta,    'product title','title','name')
    shop_title_col  = find_col(shopify, 'product title','title','name')

    if not meta_spend_col:  raise ValueError("Cannot detect spend column in Meta CSV.")
    if not shopify_rev_col: raise ValueError("Cannot detect revenue column in Shopify CSV.")

    meta    = meta.copy();    shopify = shopify.copy()
    meta['_pid']    = meta['Product ID'].apply(clean_pid)
    shopify['_pid'] = shopify['Product ID'].apply(clean_pid)

    title_map = {}
    if shop_title_col:
        for _, r in shopify.drop_duplicates('_pid').iterrows():
            title_map[r['_pid']] = r[shop_title_col]
    if meta_title_col:
        for _, r in meta.drop_duplicates('_pid').iterrows():
            title_map[r['_pid']] = r[meta_title_col]

    meta['_spend']  = pd.to_numeric(meta[meta_spend_col],    errors='coerce').fillna(0)
    shopify['_rev'] = pd.to_numeric(shopify[shopify_rev_col], errors='coerce').fillna(0)

    # Aggregate OVERALL (no month breakdown for quadrant)
    meta_g    = meta.groupby('_pid')['_spend'].sum().reset_index()
    shopify_g = shopify.groupby('_pid')['_rev'].sum().reset_index()

    merged = pd.merge(meta_g, shopify_g, on='_pid', how='outer').fillna(0)
    merged.columns = ['Product ID','Spend','Revenue']
    merged['Product Title'] = merged['Product ID'].map(title_map).fillna('Unknown')
    merged['ROI'] = (merged['Revenue'] / merged['Spend'].replace(0, float('nan'))).fillna(0).round(4)

    avg_sp = merged['Spend'].mean()
    avg_rv = merged['Revenue'].mean()
    sp_cut = avg_sp * spend_pct_thresh / 100
    rv_cut = avg_rv * rev_pct_thresh   / 100

    cols = ['Product ID','Product Title','Spend','Revenue','ROI']

    # Q1: High Revenue + Low Spend  → High Potential
    q1 = merged[(merged['Revenue']>=rv_cut) & (merged['Spend']< sp_cut)][cols].sort_values('Revenue',ascending=False).reset_index(drop=True)
    # Q2: High Revenue + High Spend → High Conversion
    q2 = merged[(merged['Revenue']>=rv_cut) & (merged['Spend']>=sp_cut)][cols].sort_values('Revenue',ascending=False).reset_index(drop=True)
    # Q3: Low Revenue  + High Spend → Low Performing
    q3 = merged[(merged['Revenue']< rv_cut) & (merged['Spend']>=sp_cut)][cols].sort_values('Spend',  ascending=False).reset_index(drop=True)
    # Q4: Low Revenue  + Low Spend  → Zombie
    q4 = merged[(merged['Revenue']< rv_cut) & (merged['Spend']< sp_cut)][cols].sort_values('Revenue',ascending=False).reset_index(drop=True)

    return {'q1':q1,'q2':q2,'q3':q3,'q4':q4,
            'all':merged,'sp_cut':sp_cut,'rv_cut':rv_cut,
            'avg_sp':avg_sp,'avg_rv':avg_rv}


# ══════════════════════════════════════════════════════════════════════
#  CHART HELPERS  (Section 1)
# ══════════════════════════════════════════════════════════════════════
def make_bar_chart(results_df, months_ordered, metric_key, metric_label, fmt='inr'):
    disc_vals  = [results_df[(results_df['Month']==m)&(results_df['Category']=='Discounted')][metric_key].iloc[0]
                  for m in months_ordered]
    ndisc_vals = [results_df[(results_df['Month']==m)&(results_df['Category']=='Non-Discounted')][metric_key].iloc[0]
                  for m in months_ordered]

    if fmt=='pct':
        text_d = [f"{v*100:.1f}%" for v in disc_vals]
        text_n = [f"{v*100:.1f}%" for v in ndisc_vals]
    elif fmt=='roi':
        text_d = [f"{v:.2f}x" for v in disc_vals]
        text_n = [f"{v:.2f}x" for v in ndisc_vals]
    else:
        text_d = [f"₹{v:,.0f}" for v in disc_vals]
        text_n = [f"₹{v:,.0f}" for v in ndisc_vals]

    fig = go.Figure()
    fig.add_trace(go.Bar(name='Discounted',     x=months_ordered, y=disc_vals,
                         marker_color='#3B82F6', text=text_d,
                         textposition='outside', textfont=dict(size=11,color='white')))
    fig.add_trace(go.Bar(name='Non-Discounted', x=months_ordered, y=ndisc_vals,
                         marker_color='#10B981', text=text_n,
                         textposition='outside', textfont=dict(size=11,color='white')))
    fig.update_layout(
        title=dict(text=metric_label, font=dict(size=14,color='white'), x=0),
        barmode='group',
        plot_bgcolor='#0E1117', paper_bgcolor='#0E1117',
        font=dict(color='white'),
        xaxis=dict(showgrid=False, tickfont=dict(color='white',size=11)),
        yaxis=dict(showgrid=True,  gridcolor='#2D2D3A', tickfont=dict(color='white')),
        legend=dict(font=dict(color='white'), bgcolor='rgba(0,0,0,0)'),
        margin=dict(t=50,b=30,l=10,r=10),
        height=320,
    )
    return fig

def make_roi_line(results_df, months_ordered):
    disc_roi  = [results_df[(results_df['Month']==m)&(results_df['Category']=='Discounted')]['ROI'].iloc[0]
                 for m in months_ordered]
    ndisc_roi = [results_df[(results_df['Month']==m)&(results_df['Category']=='Non-Discounted')]['ROI'].iloc[0]
                 for m in months_ordered]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=months_ordered,y=disc_roi,  mode='lines+markers+text',
                             name='Discounted',     line=dict(color='#3B82F6',width=3),
                             marker=dict(size=9),
                             text=[f"{v:.2f}x" for v in disc_roi],
                             textposition='top center',textfont=dict(color='#93C5FD',size=11)))
    fig.add_trace(go.Scatter(x=months_ordered,y=ndisc_roi, mode='lines+markers+text',
                             name='Non-Discounted', line=dict(color='#10B981',width=3),
                             marker=dict(size=9),
                             text=[f"{v:.2f}x" for v in ndisc_roi],
                             textposition='top center',textfont=dict(color='#6EE7B7',size=11)))
    fig.update_layout(
        title=dict(text='ROI Trend by Month', font=dict(size=14,color='white'),x=0),
        plot_bgcolor='#0E1117', paper_bgcolor='#0E1117',
        font=dict(color='white'),
        xaxis=dict(showgrid=False,tickfont=dict(color='white',size=11)),
        yaxis=dict(showgrid=True,gridcolor='#2D2D3A',tickfont=dict(color='white'),
                   ticksuffix='x'),
        legend=dict(font=dict(color='white'),bgcolor='rgba(0,0,0,0)'),
        margin=dict(t=50,b=30,l=10,r=10), height=320)
    return fig


# ══════════════════════════════════════════════════════════════════════
#  QUADRANT SCATTER CHART  (Section 2)
# ══════════════════════════════════════════════════════════════════════
def make_quadrant_chart(data, sp_cut, rv_cut):
    all_df = data['all'].copy()
    all_df['Quadrant'] = 'Zombie'
    all_df.loc[(all_df['Revenue']>=rv_cut)&(all_df['Spend']< sp_cut), 'Quadrant'] = 'High Potential'
    all_df.loc[(all_df['Revenue']>=rv_cut)&(all_df['Spend']>=sp_cut), 'Quadrant'] = 'High Conversion'
    all_df.loc[(all_df['Revenue']< rv_cut)&(all_df['Spend']>=sp_cut), 'Quadrant'] = 'Low Performing'

    color_map = {
        'High Potential':  '#10B981',
        'High Conversion': '#3B82F6',
        'Low Performing':  '#EF4444',
        'Zombie':          '#6B7280',
    }

    fig = go.Figure()
    for quad, color in color_map.items():
        sub = all_df[all_df['Quadrant']==quad]
        if sub.empty: continue
        fig.add_trace(go.Scatter(
            x=sub['Spend'], y=sub['Revenue'],
            mode='markers', name=quad,
            marker=dict(color=color, size=7, opacity=0.75,
                        line=dict(width=0.5, color='white')),
            text=sub['Product Title'],
            customdata=sub[['Product ID','Spend','Revenue','ROI']].values,
            hovertemplate=(
                "<b>%{text}</b><br>"
                "PID: %{customdata[0]}<br>"
                "Spend: ₹%{customdata[1]:,.0f}<br>"
                "Revenue: ₹%{customdata[2]:,.0f}<br>"
                "ROI: %{customdata[3]:.2f}x<extra></extra>"
            )
        ))

    # Threshold lines
    fig.add_vline(x=sp_cut, line_dash='dash', line_color='#FBBF24', line_width=1.5,
                  annotation_text=f"Spend cut-off ₹{sp_cut:,.0f}",
                  annotation_font=dict(color='#FBBF24',size=10),
                  annotation_position='top right')
    fig.add_hline(y=rv_cut, line_dash='dash', line_color='#FBBF24', line_width=1.5,
                  annotation_text=f"Revenue cut-off ₹{rv_cut:,.0f}",
                  annotation_font=dict(color='#FBBF24',size=10),
                  annotation_position='top right')

    # Quadrant labels (watermark style)
    max_sp = all_df['Spend'].max()
    max_rv = all_df['Revenue'].max()
    annotations = [
        dict(x=sp_cut/2,        y=max_rv*0.92, text="HIGH POTENTIAL",   showarrow=False,
             font=dict(color='#10B981',size=11,family='Arial Black'), opacity=0.4),
        dict(x=sp_cut+(max_sp-sp_cut)/2, y=max_rv*0.92, text="HIGH CONVERSION", showarrow=False,
             font=dict(color='#3B82F6',size=11,family='Arial Black'), opacity=0.4),
        dict(x=sp_cut+(max_sp-sp_cut)/2, y=rv_cut*0.5,  text="LOW PERFORMING",  showarrow=False,
             font=dict(color='#EF4444',size=11,family='Arial Black'), opacity=0.4),
        dict(x=sp_cut/2,        y=rv_cut*0.5,  text="ZOMBIE",          showarrow=False,
             font=dict(color='#9CA3AF',size=11,family='Arial Black'), opacity=0.4),
    ]

    fig.update_layout(
        title=dict(text='Product Quadrant Map — Overall Spend vs Revenue',
                   font=dict(size=15,color='white'), x=0),
        xaxis=dict(title='Total Ad Spend (INR) →', showgrid=True,
                   gridcolor='#2D2D3A', tickfont=dict(color='white'),
                   title_font=dict(color='#94A3B8')),
        yaxis=dict(title='Total Revenue (INR) →', showgrid=True,
                   gridcolor='#2D2D3A', tickfont=dict(color='white'),
                   title_font=dict(color='#94A3B8')),
        plot_bgcolor='#0E1117', paper_bgcolor='#0E1117',
        font=dict(color='white'),
        legend=dict(font=dict(color='white',size=11), bgcolor='rgba(0,0,0,0)',
                    orientation='h', yanchor='bottom', y=1.01, xanchor='left', x=0),
        annotations=annotations,
        margin=dict(t=80,b=60,l=60,r=20),
        height=540,
    )
    return fig


# ══════════════════════════════════════════════════════════════════════
#  QUADRANT CARD (Section 2 UI helper)
# ══════════════════════════════════════════════════════════════════════
def quadrant_card(label, emoji, description, color, df, key_prefix):
    n    = len(df)
    sp   = df['Spend'].sum()
    rv   = df['Revenue'].sum()
    roi  = round(rv/sp,2) if sp else 0

    border_css = f"border-left: 4px solid {color};"
    st.markdown(f"""
    <div style="background:#1E2130; border-radius:12px; padding:20px 24px; {border_css} margin-bottom:8px;">
        <div style="font-size:22px; margin-bottom:6px;">{emoji}</div>
        <div style="font-size:17px; font-weight:700; color:white; margin-bottom:4px;">{label}</div>
        <div style="font-size:12px; color:#94A3B8; margin-bottom:16px;">{description}</div>
        <div style="display:flex; gap:24px; flex-wrap:wrap;">
            <div><div style="font-size:11px;color:#64748B;text-transform:uppercase;letter-spacing:.06em">Products</div>
                 <div style="font-size:20px;font-weight:700;color:{color}">{n}</div></div>
            <div><div style="font-size:11px;color:#64748B;text-transform:uppercase;letter-spacing:.06em">Total Spend</div>
                 <div style="font-size:20px;font-weight:700;color:white">₹{sp:,.0f}</div></div>
            <div><div style="font-size:11px;color:#64748B;text-transform:uppercase;letter-spacing:.06em">Total Revenue</div>
                 <div style="font-size:20px;font-weight:700;color:white">₹{rv:,.0f}</div></div>
            <div><div style="font-size:11px;color:#64748B;text-transform:uppercase;letter-spacing:.06em">Quadrant ROI</div>
                 <div style="font-size:20px;font-weight:700;color:{color}">{roi:.2f}x</div></div>
        </div>
    </div>
    """, unsafe_allow_html=True)

    with st.expander(f"📋 View {label} product list ({n} products)"):
        if df.empty:
            st.info("No products in this quadrant.")
        else:
            disp = df.copy()
            disp['Spend']   = disp['Spend'].apply(fmt_inr)
            disp['Revenue'] = disp['Revenue'].apply(fmt_inr)
            disp['ROI']     = disp['ROI'].apply(fmt_roi)
            st.dataframe(disp, hide_index=True)


# ══════════════════════════════════════════════════════════════════════
#  SECTION 1 UI
# ══════════════════════════════════════════════════════════════════════
if page == "📋 Discount Analysis":
    st.title("📋 Discount Analysis")
    st.caption("Monthly Discounted vs Non-Discounted performance — summary matrix, charts, and product-level insights.")

    if run_s1:
        if not meta_file or not shopify_file or not discount_file:
            st.error("Please upload all 3 files to run this analysis.")
            st.stop()

        with st.spinner("Processing data…"):
            try:
                meta_df     = pd.read_csv(meta_file)
                shopify_df  = pd.read_csv(shopify_file)
                discount_df = (pd.read_csv(discount_file)
                               if discount_file.name.endswith('.csv')
                               else pd.read_excel(discount_file))
                results_df, months_ordered, merged, insights, title_map = run_discount_analysis(
                    meta_df, shopify_df, discount_df, spend_pct, rev_pct)
            except Exception as e:
                st.error(f"Error: {e}"); st.exception(e); st.stop()

        # ── top KPI bar ──────────────────────────────────────────────
        st.success(f"✅ {len(months_ordered)} month(s) detected: {', '.join(months_ordered)}")
        k1,k2,k3,k4 = st.columns(4)
        k1.metric("Total Products",      merged['Product ID'].nunique())
        k2.metric("Discounted",          merged[merged['Is_Discounted']]['Product ID'].nunique())
        k3.metric("Non-Discounted",      merged[~merged['Is_Discounted']]['Product ID'].nunique())
        total_roi = round(merged['Revenue'].sum()/merged['Spend'].sum(),2) if merged['Spend'].sum() else 0
        k4.metric("Overall ROI",         f"{total_roi:.2f}x")

        st.markdown("---")

        # ── Tabs ─────────────────────────────────────────────────────
        tab1, tab2, tab3 = st.tabs(["📊 Summary Matrix & Charts", "🔍 Product Insights", "⬇ Download"])

        # ── Tab 1: Matrix + Charts ────────────────────────────────────
        with tab1:
            st.subheader("Monthly Summary Matrix")
            for month in months_ordered:
                st.markdown(f"#### 📅 {month}")
                md   = results_df[results_df['Month']==month].copy()
                disp = md[['Category','Spend','Revenue','Spend_Pct','Revenue_Pct','ROI']].copy()
                disp.columns=['Category','Spend (INR)','Revenue (INR)','Spend %','Revenue %','ROI']
                disp['Spend (INR)']   = disp['Spend (INR)'].apply(fmt_inr)
                disp['Revenue (INR)'] = disp['Revenue (INR)'].apply(fmt_inr)
                disp['Spend %']       = disp['Spend %'].apply(fmt_pct)
                disp['Revenue %']     = disp['Revenue %'].apply(fmt_pct)
                disp['ROI']           = disp['ROI'].apply(fmt_roi)
                st.dataframe(disp.set_index('Category'))

            st.markdown("---")
            st.subheader("📈 Charts")

            # Row 1: Spend + Revenue bars
            c1, c2 = st.columns(2)
            with c1:
                st.plotly_chart(make_bar_chart(results_df, months_ordered,
                    'Spend','Total Spend by Month (INR)','inr'))
            with c2:
                st.plotly_chart(make_bar_chart(results_df, months_ordered,
                    'Revenue','Total Revenue by Month (INR)','inr'))

            # Row 2: ROI line + Spend %
            c3, c4 = st.columns(2)
            with c3:
                st.plotly_chart(make_roi_line(results_df, months_ordered))
            with c4:
                st.plotly_chart(make_bar_chart(results_df, months_ordered,
                    'Spend_Pct','Spend Share by Month (%)','pct'))

        # ── Tab 2: Product Insights ───────────────────────────────────
        with tab2:
            st.caption(f"Thresholds — High Spend: >{spend_pct}% of category monthly average · "
                       f"Low Revenue: <{rev_pct}% of category monthly average")

            BUCKET_META = {
                'hslr': ('🔴','High Spend · Low Revenue',
                         'Consuming budget with poor returns — review or pause.'),
                'lshr': ('🟢','Low Spend · High Revenue',
                         'Efficient performers — consider scaling budget.'),
            }

            for month in months_ordered:
                st.markdown(f"---\n### 📅 {month}")

                for cat in ['Discounted','Non-Discounted']:
                    ins = insights.get((month,cat),{})
                    cat_color = '#3B82F6' if cat=='Discounted' else '#10B981'
                    st.markdown(
                        f"<div style='font-size:15px;font-weight:700;color:{cat_color};"
                        f"border-left:3px solid {cat_color};padding-left:10px;"
                        f"margin:12px 0 8px'>{cat}</div>",
                        unsafe_allow_html=True)

                    col_a, col_b = st.columns(2)
                    for col_obj, itype in zip([col_a, col_b], ['hslr','lshr']):
                        icon,label,tip = BUCKET_META[itype]
                        df_sec = ins.get(itype, pd.DataFrame())
                        with col_obj:
                            st.markdown(f"**{icon} {label}**")
                            st.caption(tip)
                            if df_sec.empty:
                                st.info("No products match this criteria.")
                            else:
                                # Summary totals
                                t_sp  = df_sec['Spend'].sum()
                                t_rv  = df_sec['Revenue'].sum()
                                t_roi = round(t_rv/t_sp,2) if t_sp else 0
                                m1,m2,m3,m4 = st.columns(4)
                                m1.metric("Products", len(df_sec))
                                m2.metric("Spend",    fmt_inr(t_sp))
                                m3.metric("Revenue",  fmt_inr(t_rv))
                                m4.metric("ROI",      fmt_roi(t_roi))

                                disp = df_sec.copy()
                                disp['Spend']   = disp['Spend'].apply(fmt_inr)
                                disp['Revenue'] = disp['Revenue'].apply(fmt_inr)
                                disp['ROI']     = disp['ROI'].apply(fmt_roi)
                                disp = disp.rename(columns={'Product ID':'PID',
                                                            'Product Title':'Title',
                                                            'Spend':'Spend (INR)',
                                                            'Revenue':'Revenue (INR)'})
                                st.dataframe(disp, hide_index=True)

        # ── Tab 3: Download ───────────────────────────────────────────
        with tab3:
            st.subheader("⬇ Download Full Excel Report")
            st.markdown("""
The Excel file contains:
- **Sheet 1 — Summary Matrix**: All months × Discounted / Non-Discounted with 5 metrics
- **Sheet 2 — Product Insights**: Every month × category bucket with totals row, colour-coded
""")
            excel_buf = build_s1_excel(results_df, months_ordered, insights)
            st.download_button(
                label     = f"📥 Download {brand_name} Discount Analysis.xlsx",
                data      = excel_buf,
                file_name = f"{brand_name}_discount_analysis.xlsx",
                mime      = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type      = "primary")
    else:
        st.info("👈 Upload your 3 CSV files in the sidebar and click **Run Discount Analysis**.")
        with st.expander("ℹ️ How this section works"):
            st.markdown("""
1. **Summary Matrix** — Month-over-month view of Spend, Revenue, %, and ROI split by Discounted vs Non-Discounted
2. **Charts** — Bar charts for Spend/Revenue, ROI trend line, Spend share — all labelled with values
3. **Product Insights** — For each month and category, products bucketed into:
   - 🔴 **High Spend · Low Revenue** — draining budget, candidates to pause
   - 🟢 **Low Spend · High Revenue** — efficient, candidates to scale
   - Each bucket shows a **totals summary** (count, spend, revenue, ROI) plus the full product list
""")


# ══════════════════════════════════════════════════════════════════════
#  SECTION 2 UI
# ══════════════════════════════════════════════════════════════════════
else:
    st.title("🔲 Product Analysis")
    st.caption("Overall 4-quadrant view across all months — identify your best and worst performing products.")

    if run_s2:
        if not s2_meta_file or not s2_shopify_file:
            st.error("Please upload Meta and Shopify CSV files.")
            st.stop()

        with st.spinner("Processing data…"):
            try:
                s2_meta_df    = pd.read_csv(s2_meta_file)
                s2_shopify_df = pd.read_csv(s2_shopify_file)
                data = run_product_analysis(s2_meta_df, s2_shopify_df, s2_spend_pct, s2_rev_pct)
            except Exception as e:
                st.error(f"Error: {e}"); st.exception(e); st.stop()

        sp_cut = data['sp_cut']; rv_cut = data['rv_cut']
        all_df = data['all']

        # KPI bar
        total_sp  = all_df['Spend'].sum()
        total_rv  = all_df['Revenue'].sum()
        total_roi = round(total_rv/total_sp,2) if total_sp else 0
        k1,k2,k3,k4 = st.columns(4)
        k1.metric("Total Products",  len(all_df))
        k2.metric("Total Spend",     fmt_inr(total_sp))
        k3.metric("Total Revenue",   fmt_inr(total_rv))
        k4.metric("Overall ROI",     fmt_roi(total_roi))

        st.info(
            f"**Thresholds** — High Spend: ≥ ₹{sp_cut:,.0f} (={s2_spend_pct}% of avg ₹{data['avg_sp']:,.0f})  "
            f"·  High Revenue: ≥ ₹{rv_cut:,.0f} (={s2_rev_pct}% of avg ₹{data['avg_rv']:,.0f})")

        st.markdown("---")

        # ── Quadrant scatter chart ────────────────────────────────────
        st.plotly_chart(make_quadrant_chart(data, sp_cut, rv_cut))
        st.caption("X-axis = Total Ad Spend (INR) across all months  ·  Y-axis = Total Revenue (INR) across all months  ·  Hover a dot for product details")

        st.markdown("---")

        # ── 4 quadrant cards (2×2 grid) ──────────────────────────────
        st.subheader("Quadrant Breakdown")
        st.caption("Click 'View product list' inside any quadrant to see all products.")

        row1_c1, row1_c2 = st.columns(2)
        row2_c1, row2_c2 = st.columns(2)

        with row1_c1:
            quadrant_card(
                "High Potential Products", "🚀",
                "High Revenue · Low Spend — great ROI, underinvested",
                "#10B981", data['q1'], "q1")

        with row1_c2:
            quadrant_card(
                "High Conversion Products", "💎",
                "High Revenue · High Spend — strong performers, worth the budget",
                "#3B82F6", data['q2'], "q2")

        with row2_c1:
            quadrant_card(
                "Low Performing Products", "⚠️",
                "Low Revenue · High Spend — budget drain, review urgently",
                "#EF4444", data['q3'], "q3")

        with row2_c2:
            quadrant_card(
                "Zombie Products", "🧟",
                "Low Revenue · Low Spend — minimal activity, assess or drop",
                "#6B7280", data['q4'], "q4")

        # Axis legend note
        st.markdown("""
        <div style="background:#1E2130;border-radius:8px;padding:12px 18px;margin-top:8px;
                    font-size:12px;color:#94A3B8;">
        <b>Axis guide</b><br>
        <b>X-axis (horizontal)</b> = Total Ad Spend (INR) per product across all months uploaded<br>
        <b>Y-axis (vertical)</b> = Total Revenue (INR) per product across all months uploaded<br>
        Products to the <b>right</b> of the dashed line = High Spend ·
        Products <b>above</b> the dashed line = High Revenue
        </div>""", unsafe_allow_html=True)

    else:
        st.info("👈 Upload Meta and Shopify CSV files in the sidebar and click **Run Product Analysis**.")
        with st.expander("ℹ️ How this section works"):
            st.markdown("""
**4-Quadrant Product Map** — all months aggregated together, each product gets one overall Spend and Revenue value.

| Quadrant | Spend | Revenue | What it means |
|---|---|---|---|
| 🚀 High Potential | Low | High | Great ROI — increase budget |
| 💎 High Conversion | High | High | Strong performers — maintain |
| ⚠️ Low Performing | High | Low | Budget drain — pause/review |
| 🧟 Zombie | Low | Low | Minimal activity — assess or drop |

**Thresholds** are set as a % of the overall average spend/revenue per product.
For example, 100% means the average — a product above 100% of average spend is "High Spend".
""")